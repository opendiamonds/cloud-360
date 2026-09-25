"""Estimate-table readers (CSV / XLSX) — no HTTP/DB/framework imports."""

from __future__ import annotations

import csv
import io
import zipfile
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from typing import Iterable, Sequence

MAX_PAYLOAD_BYTES = 32 * 1024 * 1024
MAX_ROWS = 50_000
HEADER_SCAN_ROWS = 40

# Distinctive header signatures for cloud detection (case-insensitive match).
# Include aliases seen in official Pricing Calculator exports.
AWS_HEADERS = frozenset(
    {
        "description",
        "service",
        "region",
        "quantity",
        "monthly cost",
        "monthly",
        "currency",
        "instance type",
        "upfront",
        "first year total",
        "first 12 months total",
        "group hierarchy",
        "configuration summary",
    }
)
GCP_HEADERS = frozenset(
    {
        # Legacy Pricing Calculator CSV
        "sku description",
        "sku id",
        "quantity",
        "unit",
        "cost",
        "currency",
        "service description",
        "usage",
        # Current Google Cloud Pricing Calculator export
        "service_display_name",
        "service_id",
        "sku",
        "total_price",
        "total_price, usd",
        "name",
    }
)
AZURE_HEADERS = frozenset(
    {
        "service name",
        "service category",
        "service type",
        "region",
        "quantity",
        "estimated monthly cost",
        "estimated upfront cost",
        "currency",
        "custom name",
        "description",
    }
)

_CLOUD_HEADER_SETS = {
    "aws": AWS_HEADERS,
    "gcp": GCP_HEADERS,
    "azure": AZURE_HEADERS,
}


@dataclass(frozen=True)
class RawTable:
    headers: list[str]
    rows: list[list[str]]
    source_format: str  # csv | xlsx


def normalize_header(cell: str) -> str:
    return " ".join(cell.strip().lower().split())


def parse_number(raw: str | None) -> Decimal | None:
    if raw is None:
        return None
    text = str(raw).strip()
    if not text:
        return None
    cleaned = (
        text.replace(",", "")
        .replace("$", "")
        .replace("USD", "")
        .replace("NT$", "")
        .strip()
    )
    if not cleaned:
        return None
    try:
        value = Decimal(cleaned)
    except (InvalidOperation, ValueError):
        return None
    if not value.is_finite():
        return None
    return value


def xlsx_uncompressed_bytes(file_bytes: bytes) -> int:
    """Sum of ZIP entry uncompressed sizes (NFR3.1)."""
    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
            return sum(info.file_size for info in zf.infolist())
    except zipfile.BadZipFile:
        return -1


def _header_matches(normalized_header: str, alias: str) -> bool:
    """Exact match, or official export compound headers (e.g. Region / Location…)."""
    if not normalized_header or not alias:
        return False
    if normalized_header == alias:
        return True
    if normalized_header.startswith(alias + " /"):
        return True
    if normalized_header.startswith(alias + " ("):
        return True
    # GCP: "total_price, USD" / "total_price,USD"
    if normalized_header.startswith(alias + ","):
        return True
    return False


def _normalized_hits(headers: Sequence[str], signatures: frozenset[str]) -> int:
    norms = [normalize_header(h) for h in headers if h and str(h).strip()]
    hits = 0
    for sig in signatures:
        if any(_header_matches(n, sig) for n in norms):
            hits += 1
    return hits


def _table_from_matrix(matrix: list[list[str]], source_format: str) -> RawTable | None:
    """Pick the row with the best cloud-signature score as headers (title rows skipped)."""
    if not matrix:
        return None
    best_idx = 0
    best_key = (-1, -1)
    scan = min(len(matrix), HEADER_SCAN_ROWS)
    for i in range(scan):
        row = matrix[i]
        if not any(str(c).strip() for c in row):
            continue
        score = max(score_cloud(row).values())
        density = sum(1 for c in row if str(c).strip())
        key = (score, density)
        if key > best_key:
            best_key = key
            best_idx = i
    headers = [str(c) for c in matrix[best_idx]]
    body = [[str(c) for c in row] for row in matrix[best_idx + 1 :]]
    return RawTable(headers=headers, rows=body, source_format=source_format)


def read_csv_table(file_bytes: bytes) -> RawTable | None:
    try:
        text = file_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = file_bytes.decode("latin-1")
        except UnicodeDecodeError:
            return None
    if len(text.encode("utf-8")) > MAX_PAYLOAD_BYTES:
        return None  # caller maps to ambiguous
    try:
        reader = csv.reader(io.StringIO(text))
        rows = [list(r) for r in reader]
    except csv.Error:
        return None
    if not rows:
        return None
    return _table_from_matrix([[str(c) for c in r] for r in rows], "csv")


def read_xlsx_table(file_bytes: bytes) -> RawTable | None:
    uncompressed = xlsx_uncompressed_bytes(file_bytes)
    if uncompressed < 0 or uncompressed > MAX_PAYLOAD_BYTES:
        return None
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ValueError("openpyxl is required to read xlsx estimates") from exc
    try:
        wb = load_workbook(
            io.BytesIO(file_bytes), read_only=True, data_only=True
        )
    except Exception:
        return None
    try:
        best: RawTable | None = None
        best_score = -1
        for ws in wb.worksheets:
            matrix: list[list[str]] = []
            for row in ws.iter_rows(values_only=True):
                matrix.append(["" if c is None else str(c) for c in row])
                if len(matrix) > MAX_ROWS + HEADER_SCAN_ROWS:
                    break
            if not matrix:
                continue
            table = _table_from_matrix(matrix, "xlsx")
            if table is None:
                continue
            score = max(score_cloud(table.headers).values())
            if score > best_score:
                best_score = score
                best = table
        return best
    finally:
        wb.close()


def sniff_format(file_bytes: bytes, filename: str | None) -> str | None:
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        return "xlsx"
    if name.endswith(".csv"):
        return "csv"
    if file_bytes[:2] == b"PK":
        return "xlsx"
    # Heuristic: printable / has commas → csv
    sample = file_bytes[:512]
    if b"," in sample or b";" in sample:
        return "csv"
    return None


def score_cloud(headers: Sequence[str]) -> dict[str, int]:
    return {
        cloud: _normalized_hits(headers, sigs)
        for cloud, sigs in _CLOUD_HEADER_SETS.items()
    }


def best_clouds(scores: dict[str, int], minimum: int = 3) -> list[str]:
    ranked = sorted(
        ((cloud, score) for cloud, score in scores.items() if score >= minimum),
        key=lambda item: (-item[1], item[0]),
    )
    if not ranked:
        return []
    top = ranked[0][1]
    return [cloud for cloud, score in ranked if score == top]


def _find_column(norm: Sequence[str], names: tuple[str, ...]) -> int | None:
    for name in names:
        for i, header in enumerate(norm):
            if _header_matches(header, name):
                return i
    return None


def column_map(headers: Sequence[str], cloud: str) -> dict[str, int]:
    """Map logical fields to column indices for a resolved cloud."""
    norm = [normalize_header(h) for h in headers]
    aliases: dict[str, dict[str, tuple[str, ...]]] = {
        "aws": {
            "itemName": ("description", "service"),
            "itemNameAlt": ("service",),
            "spec": (
                "sku",
                "sku id",
                "instance type",
                "configuration summary",
                "region",
                "service",
            ),
            "quantity": ("quantity",),
            "amount": (
                "monthly cost",
                "monthly",
                "cost",
                "amount",
                "first 12 months total",
                "first year total",
            ),
            "currency": ("currency",),
        },
        "gcp": {
            "itemName": (
                "sku description",
                "service description",
                "name",
                "service_display_name",
                "description",
            ),
            "spec": (
                "sku id",
                "sku",
                "unit",
                "region",
                "service_display_name",
            ),
            "serviceId": ("service_id", "service id"),
            "quantity": ("quantity", "usage"),
            "amount": (
                "cost",
                "amount",
                "total_price",
                "total_price, usd",
                "total price",
            ),
            "currency": ("currency", "currency code"),
        },
        "azure": {
            "itemName": (
                "service name",
                "service type",
                "custom name",
                "description",
                "service category",
            ),
            "spec": (
                "sku",
                "sku id",
                "sku name",
                "arm sku name",
                "product name",
                "service tier",
                "region",
                "service category",
            ),
            "quantity": ("quantity",),
            "amount": (
                "estimated monthly cost",
                "monthly cost",
                "cost",
                "amount",
            ),
            "currency": ("currency", "currency code"),
        },
    }
    mapping: dict[str, int] = {}
    for field, names in aliases[cloud].items():
        idx = _find_column(norm, names)
        if idx is not None:
            mapping[field] = idx
    return mapping


def row_get(row: Sequence[str], index: int | None) -> str:
    if index is None or index < 0 or index >= len(row):
        return ""
    return row[index]


def infer_currency_from_headers(
    headers: Sequence[str], amount_index: int | None
) -> str | None:
    """When exports embed currency in the amount header (e.g. total_price, USD)."""
    if amount_index is None or amount_index < 0 or amount_index >= len(headers):
        return None
    text = normalize_header(str(headers[amount_index]))
    for code in ("usd", "twd", "eur", "jpy", "gbp", "cny", "hkd"):
        if code in text.replace("_", " ").replace(",", " ").split() or code in text:
            # Prefer token / substring match for ", USD" style headers
            if code in text:
                return code.upper()
    return None


def iter_limited_rows(rows: Iterable[list[str]]) -> list[list[str]] | None:
    out: list[list[str]] = []
    for row in rows:
        out.append(row)
        if len(out) > MAX_ROWS:
            return None
    return out
